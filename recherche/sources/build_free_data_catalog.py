"""Assemble the reviewed research catalogues into a French workbook and evidence JSON.

The input files are read afresh at each invocation. This script performs no
network access and does not modify the stock bot or any source catalogue.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


BASE = Path(__file__).resolve().parent
CATALOGUES = ["root_catalog.json", "market_catalog.json", "regulators_catalog.json", "official_us_macro_catalog.json"]
TEST_FILES = ["official_us_macro_tests.json", "regulators_tests.json", "root_tests.json", "market_tests.json", "market_quote_tests.json"]
NAVY = "123344"
TEAL = "177B78"
INK = "203C49"
MUTED = "617983"
PALE = "EAF4F4"
WHITE = "FFFFFF"
LINE = "D7E3E8"
AMBER = "FFF1D4"
GREEN = "E3F3E9"
GREY = "EFF2F4"

EVIDENCE_LABELS = {
    "documentation_verified": "Documentation vérifiée",
    "live_tested": "Accès public testé — périmètre limité",
    "needs_account_validation": "Compte / droits d’accès à valider",
}


def safe_text(value):
    if value is None:
        return "Non renseigné"
    if isinstance(value, bool):
        return "Oui" if value else "Non"
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            return "Non renseigné"
        return value
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, allow_nan=False)
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))
    # All catalogue/test text is data, never an executable spreadsheet formula.
    if value.lstrip().startswith(("=", "+", "-", "@")):
        value = "'" + value
    if len(value) > 32700:
        raise ValueError("Une cellule dépasse la taille supportée ; conserver ce détail dans les preuves JSON.")
    return value


def safe_url(value):
    if not isinstance(value, str) or any(c.isspace() or ord(c) < 32 for c in value):
        raise ValueError(f"URL invalide dans le catalogue : {value!r}")
    if any(c in value for c in ('\\', '<', '>', '"', "'")):
        raise ValueError(f"URL invalide dans le catalogue : {value!r}")
    parts = urlsplit(value)
    if parts.scheme != "https" or not parts.hostname or parts.username or parts.password:
        raise ValueError(f"URL non HTTPS ou contenant des identifiants : {value!r}")
    _ = parts.port
    return value


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8-sig"))


def load_sources(source_dir):
    sources = []
    for filename in CATALOGUES:
        path = source_dir / filename
        if not path.is_file():
            raise FileNotFoundError(f"Catalogue requis encore absent : {path}")
        data = read_json(path)
        if not isinstance(data, list):
            raise ValueError(f"{filename} doit contenir une liste.")
        for original in data:
            item = dict(original)
            item["catalogue_file"] = filename
            item["priority_original"] = item.get("priority")
            item["assembly_note"] = ""
            name = str(item.get("source", ""))
            if not name:
                raise ValueError(f"Source sans nom dans {filename}")
            if "yfinance" in name.casefold():
                item["priority"] = "B"
                item["assembly_note"] = (
                    "Priorité B à l’assemblage : accès Yahoo non officiel et fragile. "
                    "Quatre exercices peuvent satisfaire le minimum de trois du moteur si les champs sont complets et cohérents."
                )
            elif "fiscal.ai" in name.casefold():
                item["priority"] = "B"
                item["assembly_note"] = (
                    "Priorité B à l’assemblage : échantillon de sociétés imposé et conditions d’essai à confirmer ; "
                    "aucune couverture mondiale gratuite validée."
                )
            if item.get("priority") not in {"A", "B", "C"}:
                raise ValueError(f"Priorité invalide : {name}")
            if item.get("evidence_status") not in EVIDENCE_LABELS:
                raise ValueError(f"Niveau de preuve inconnu : {name}")
            item["urls"] = list(dict.fromkeys(safe_url(url) for url in item.get("urls", [])))
            item["source_id"] = f"S{len(sources) + 1:03d}"
            sources.append(item)
    names = [item["source"].casefold() for item in sources]
    if len(names) != len(set(names)):
        raise ValueError("Doublon exact de source : une décision de fusion est nécessaire.")
    return sources


def human_test(row, batch, filename):
    name = row.get("source") or row.get("name") or "Source non renseignée"
    known_names = {
        "sec_companyfacts_apple": "SEC EDGAR — Companyfacts Apple",
        "sec_submissions_apple": "SEC EDGAR — Submissions Apple",
        "nasdaq_symbols": "Nasdaq Trader — Symbol Directory",
        "ecb_usd_chf": "BCE — changes USD/EUR et CHF/EUR",
    }
    source = known_names.get(name, name)
    status = row.get("http_status", row.get("status"))
    timestamp = (row.get("tested_at_utc") or row.get("tested_at") or row.get("tested_on")
                 or row.get("date_utc") or batch.get("timestamp") or "Date non renseignée")
    level = "Retour à examiner"
    quality = "Validation métier non établie"
    scope = row.get("test_scope") or row.get("test_type") or "Requête publique ponctuelle"
    observations = row.get("observations", {})
    summary = row.get("result") or row.get("note") or "Voir les preuves JSON."
    limits = row.get("scope_limit") or row.get("http_status_note") or "Un exemple ne démontre pas la couverture du fournisseur."
    documented_response = False
    returned_data_or_metadata = False
    if name == "sec_companyfacts_apple":
        level = "Faits financiers bruts reçus"
        quality = "Exemples contrôlés ; normalisation complète à construire"
        scope = "Une entreprise — Apple"
        summary = (
            f"{observations.get('tags', {}).get('us-gaap', '?')} concepts US-GAAP et "
            f"{observations.get('tags', {}).get('dei', '?')} DEI. "
            "Différences vérifiées entre actions fin d’exercice, DEI à date de couverture et moyenne diluée ; "
            "dividendes ordinaires déclarés présents malgré un tag spécifique de paiement ancien."
        )
        limits = (
            "Les périodes start/end, classes et unités restent à normaliser. fy ne suffit pas. "
            "Tag absent ≠ zéro ; données standard de l’entité entière, extensions/dimensions possibles hors API."
        )
        returned_data_or_metadata = True
    elif name == "sec_submissions_apple":
        level = "Métadonnées de dépôts reçues"
        quality = "Structure et historique annoncé contrôlés"
        scope = "Une entreprise — Apple"
        summary = (
            f"{observations.get('recent_filing_count', '?')} dépôts récents ; "
            "fichier complémentaire d’historique annoncé. Ticker et place disponibles."
        )
        limits = "Liste de dépôts et dates : aucun jeu supplémentaire de comptes financiers n’a été analysé par cet appel."
        returned_data_or_metadata = True
    elif name == "nasdaq_symbols":
        level = "Répertoire de titres reçu"
        quality = "Lignes et drapeaux contrôlés ; filtrage métier à construire"
        scope = "Un fichier Nasdaq courant"
        summary = (
            f"{observations.get('rows', '?')} lignes, dont {observations.get('etf_count', '?')} ETF "
            f"et {observations.get('test_issue_count', '?')} titres de test."
        )
        limits = "Le total n’est pas un nombre d’actions ordinaires. Aucun cours, historique de composition ou fondamental fourni."
        returned_data_or_metadata = True
    elif name == "ecb_usd_chf":
        level = "Observations de change reçues"
        quality = "Unités et dates contrôlées sur quatre observations"
        scope = "Deux devises, deux observations par série"
        summary = "Quatre observations USD/EUR et CHF/EUR ; les valeurs expriment des unités de devise pour 1 EUR."
        limits = "Taux de référence : aucun prix d’action ni cours d’exécution ; ne pas inverser le sens du change."
        returned_data_or_metadata = True
    elif "swagger" in str(name).casefold():
        level = "Documentation technique reçue"
        quality = "Aucun état financier par cet appel"
        scope = "Schéma de l’API officielle TWSE"
        summary = f"Schéma Swagger reçu : {row.get('path_count', '?')} routes documentées, dont plusieurs routes de fondamentaux."
        limits = "Un schéma HTTP 200 ne prouve pas le retour de données financières ; un test distinct des comptes a été effectué."
        documented_response = True
    elif name == "filings.xbrl.org API list":
        level = "Métadonnées de rapport reçues"
        quality = "Découverte d’un rapport ; comptes non validés par cet appel"
        scope = "Un rapport danois dans la liste API"
        summary = f"{row.get('records_returned', '?')} résultat ; période finissant le {row.get('selected_period_end', '?')}."
        limits = "Un résultat de liste n’établit ni couverture par pays ni validité des chiffres du rapport."
        returned_data_or_metadata = True
    elif name == "filings.xbrl.org facts":
        level = "Faits XBRL reçus — erreurs signalées"
        quality = "Validation requise : 45 erreurs signalées à la source"
        scope = "Un rapport danois récent"
        summary = f"{row.get('facts_count', '?')} faits avec valeur, décimales et dimensions ; 45 erreurs et 1 avertissement signalés par l’API source."
        limits = "Les chiffres n’ont pas été déclarés fiables. Traiter dimensions, comptes consolidés/séparés et erreurs de validation avant calcul."
        returned_data_or_metadata = True
    elif name == "Info-financiere one record":
        level = "Métadonnées de document reçues"
        quality = "Aucun fichier financier sous-jacent testé"
        scope = "Un document réglementaire"
        summary = f"{row.get('records_returned', '?')} résultat dans un catalogue annoncé de {row.get('total_count', '?')} entrées ; liens de récupération et identifiants présents."
        limits = "Le total désigne des entrées documentaires, pas des sociétés. Le contenu financier lié n’a pas été téléchargé."
        returned_data_or_metadata = True
    elif name == "TWSE listed-company income statements":
        level = "Comptes de résultat bruts reçus"
        quality = "Retour et période contrôlés ; unités et cumul à normaliser"
        scope = "Instantané des sociétés ordinaires concernées — 2026 T2"
        summary = f"{row.get('records_returned', '?')} lignes ; période locale 115-2, soit 2026 T2. Une seule période reçue."
        limits = "Aucun historique démontré. Traduire les champs et vérifier calendrier local, unités et nature trimestrielle ou cumulée."
        returned_data_or_metadata = True
    elif "nestlé" in str(name).casefold():
        level = "Document financier XLSX lu"
        quality = "Fichier et ligne de chiffre d’affaires lus"
        scope = row.get("test_scope") or "Un document officiel Nestlé"
        worksheets = row.get("worksheets", [])
        sheet = worksheets[0] if worksheets else {}
        summary = f"Classeur lisible : {sheet.get('rows', '?')} lignes et {sheet.get('columns', '?')} colonnes ; ligne Sales extraite. Empreinte SHA-256 conservée."
        limits = "Un compte de résultat d’un émetteur ; cela ne fournit pas à lui seul tous les états ni un connecteur générique suisse."
        returned_data_or_metadata = True
    elif "précontrôle" in str(name).casefold() and "yahoo" in str(name).casefold():
        level = "Aucune donnée — précontrôle Yahoo bloqué"
        quality = "Aucun appel chart ; yfinance non validé"
        scope = "Précontrôle de la page publique NESN.SW uniquement"
        summary = "La page du ticker a renvoyé HTTP 429 via l’outil web. Arrêt immédiat ; aucun appel à l’endpoint chart, aucun cours reçu."
        limits = "Ce refus de page ne prouve pas le résultat de l’API chart ni de yfinance. Heure exacte non capturée ; horodatage d’enregistrement distinct dans le JSON."
    elif row.get("usable_market_data") is False:
        level = "Aucune donnée de marché reçue"
        quality = "Accès aux chiffres non validé"
    return {
        "source": source, "test_scope": scope, "timestamp": timestamp,
        "http_status": status, "http_status_note": row.get("http_status_note"),
        "response_level": level, "business_validation": quality,
        "summary": summary, "limitations": limits,
        "url": safe_url(row["url"]), "bytes": row.get("bytes"),
        "input_file": filename,
        "document_only": documented_response,
        "returned_data_or_metadata": returned_data_or_metadata,
    }


def load_tests(source_dir):
    originals, normalized, manifests = {}, [], []
    for filename in TEST_FILES:
        path = source_dir / filename
        payload = read_json(path)
        originals[filename] = payload
        batch = payload if isinstance(payload, dict) else {}
        records = payload.get("results", []) if isinstance(payload, dict) else payload
        if not isinstance(records, list):
            raise ValueError(f"Format de tests incompatible : {filename}")
        manifests.append({"filename": filename, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "observations": len(records)})
        for record in records:
            item = human_test(record, batch, filename)
            item["test_id"] = f"T{len(normalized) + 1:03d}"
            normalized.append(item)
    return originals, normalized, manifests


def prepare_sheet(wb, title, subtitle, columns, widths, freeze="C6"):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85 if title == "Synthese" else 75
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.sheet_properties.tabColor = NAVY if title == "Synthese" else TEAL
    last = len(columns)
    # Merges are limited to the title/subtitle area, never data cells.
    for row in (1, 2, 3):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
    ws.cell(1, 1, safe_text(title if title != "Synthese" else "DONNÉES BOURSIÈRES GRATUITES"))
    ws.cell(2, 1, safe_text(subtitle))
    ws.cell(3, 1, "Recherche au 6 octobre 2026 · Fiches de sources et d’outils · Collecteurs gratuits non intégrés par cette recherche")
    for row in (1, 2, 3):
        for col in range(1, last + 1):
            cell = ws.cell(row, col)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(name="Calibri", size=19 if row == 1 else 11, bold=row == 1, color=WHITE if row != 3 else "C7DFE5")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 35
    ws.row_dimensions[2].height = 37
    ws.row_dimensions[3].height = 25
    ws.row_dimensions[4].height = 10
    for col, (label, width) in enumerate(zip(columns, widths), 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = width
        cell = ws.cell(5, col, label)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[5].height = 32
    ws.freeze_panes = freeze
    ws.print_title_rows = "1:5"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.35, bottom=0.4, header=0.15, footer=0.15)
    ws.oddFooter.center.text = "Données gratuites · Recherche au 06.10.2026 · Page &P / &N"
    ws.oddFooter.center.size = 9
    ws.oddFooter.center.color = MUTED
    return ws


def append_rows(ws, rows, table_name=None):
    for values in rows:
        number = ws.max_row + 1
        max_lines = 1
        for col, value in enumerate(values, 1):
            cell = ws.cell(number, col, safe_text(value))
            cell.font = Font(name="Calibri", size=10, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = PatternFill("solid", fgColor=WHITE if number % 2 == 0 else "F3F8F9")
            cell.border = Border(bottom=Side(style="hair", color=LINE))
            width = ws.column_dimensions[get_column_letter(col)].width
            text = str(cell.value)
            lines = sum(max(1, math.ceil(len(part) / max(8, width - 5))) for part in text.splitlines())
            max_lines = max(max_lines, lines)
        ws.row_dimensions[number].height = min(350, max(32, max_lines * 13 + 11))
    if table_name and ws.max_row >= 6:
        table = Table(displayName=table_name, ref=f"A5:{get_column_letter(ws.max_column)}{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)
        ws.auto_filter.ref = table.ref
    ws.print_area = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"


def format_priority(cell):
    colors = {"A": ("DDF1ED", "146456"), "B": ("E8F0FA", "2E537C"), "C": (GREY, "65747D")}
    fill, foreground = colors.get(str(cell.value), (GREY, INK))
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.font = Font(name="Calibri", size=11, bold=True, color=foreground)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def link_cell(cell, url):
    cell.hyperlink = safe_url(url)
    cell.font = Font(name="Calibri", size=10, underline="single", color="176C9A")
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_summary(wb, sources, tests, stats):
    ws = prepare_sheet(
        wb, "Synthese",
        "Une architecture sans abonnement de données est possible sur un univers défini ; sa couverture et sa qualité doivent être mesurées.",
        ["Rubrique", "Constat / décision", "Comment l’interpréter", "État / repère"],
        [25, 58, 96, 43], freeze="A6",
    )
    priorities = stats["priority_counts"]
    rows = [
        ["Périmètre", f"{len(sources)} fiches de sources et d’outils", "Ces fiches incluent des régulateurs, agrégateurs, outils, échantillons et exemples d’émetteurs. Ce ne sont pas autant de fournisseurs ou de flux indépendants.", "Catalogue de recherche"],
        ["Accès gratuits", "0 abonnement est possible pour un socle ciblé", "Assembler dépôts officiels, référentiels et sources complémentaires. Les univers autorisés, historiques et droits varient ; certains services listés sont des essais ou des pistes non retenues pour une collecte automatique.", "Aucune couverture mondiale exhaustive validée"],
        ["Recherche réalisée", f"{len(tests)} observations documentées", "Les preuves intégrales et URLs sont conservées dans Tests_acces_donnees_gratuites.json. Les heures absentes ou codes HTTP non exposés restent signalés comme tels.", "Onglet Tests"],
        ["Transport HTTP", f"{stats['http_200_count']} réponses HTTP 200 mesurées", f"{stats['returned_data_or_metadata_count']} retours de données ou métadonnées et {stats['document_only_count']} schéma documentaire. Un HTTP 200 ne signifie ni chiffres validés ni droits confirmés sur toute une offre.", "Validation métier séparée"],
        ["Sans chiffres reçus", f"{stats['no_data_observations_count']} observations sans données de marché", "Exemple Alpha Vantage informatif et deux pages Stooq : codes HTTP non exposés. Précontrôle d’une page Yahoo : HTTP 429, arrêt immédiat, aucun appel chart. Ces observations ne valident ni l’API chart ni yfinance.", "3 statuts non exposés + 1 refus 429"],
        ["État du bot", "Collecteurs gratuits pas encore intégrés", "Cette recherche a vérifié documentation et exemples d’accès. Elle n’a pas remplacé le fournisseur de l’ancien bot, activé un service permanent ou validé un scan mondial complet.", "Intégration à construire"],
        ["Priorité A", f"{priorities.get('A', 0)} pistes structurantes", "À traiter d’abord pour l’architecture ou un marché ciblé. A désigne une priorité de travail, pas une garantie de qualité, de gratuité illimitée ou de profondeur suffisante.", "Validation des droits et des données requise"],
        ["Priorité B", f"{priorities.get('B', 0)} compléments", "Apport utile sur un besoin limité, une petite liste ou un pays. Inclut yfinance (accès non officiel) et Fiscal.ai (échantillon imposé/essai) rétrogradés à B pour l’assemblage.", "Complément, secours ou expérimentation"],
        ["Priorité C", f"{priorities.get('C', 0)} pistes secondaires", "Veille, recherche, consultation ou solutions dont l’automatisation gratuite n’est pas établie. L’inscription au catalogue n’est pas une autorisation de collecte.", "Facultatif / conditions à examiner"],
        ["Preuve : documentation", "Documentation vérifiée", "Page officielle ou code du projet consulté. Cela ne valide pas les droits réellement accordés à une clé gratuite ni le contenu financier reçu.", "Filtre Preuve dans Sources"],
        ["Preuve : accès", "Accès public testé — périmètre limité", "Une réponse a été examinée sur l’endpoint indiqué. Le résultat peut être une liste, un schéma, des métadonnées ou des faits bruts demandant des contrôles.", "Voir portée précise dans Tests"],
        ["Preuve : compte", "Compte / droits d’accès à valider", "Service intéressant mais utilisation réelle non confirmée avec une clé ou un compte personnel. Aucun compte n’a été créé dans cette recherche.", "Activation ultérieure par l’utilisateur"],
        ["Historique du moteur", "Minimum par défaut : 3 exercices annuels", "Le moteur existant exige 3 périodes annuelles ; le fournisseur conserve au maximum 6 années. yfinance peut fournir jusqu’à 4 années, donc parfois suffire en profondeur ; ROIC.ai à 2 années ne suffit pas seul. Plus d’historique reste utile.", "À ne pas confondre : minimum 3 / maximum conservé 6"],
        ["Périmètres étroits", "Échantillon, instantané et univers sont distincts", "Fiscal.ai : sociétés imposées ; certaines offres FMP : titres limités ; TWSE testé : une seule période 2026 T2. Ni un nombre d’endpoints ni une grande liste de symboles ne prouve une couverture financière complète.", "Couverture à mesurer par société et champ"],
        ["Qualité des chiffres", "Inconnu doit rester inconnu", "Absence de dividende ou de tag ≠ zéro. Distinguer actions de fin d’exercice, nombre DEI à date de couverture et moyenne diluée. Conserver devises, périodes, unités, classes et comptes consolidés/séparés.", "Abstention si données insuffisantes"],
        ["Dates et révisions", "Conserver période, publication et collecte", "Utiliser la date publique du dépôt pour les simulations historiques, préserver amendements et versions, et ne pas utiliser rétroactivement des données révisées. ALFRED aide pour les séries macro.", "Réduire le biais d’information future"],
        ["Droits et robustesse", "Gratuit ne signifie pas sans conditions", "Vérifier quotas, usage personnel/commercial, attribution et robots.txt quand pertinent. Les licences du code et des données sont distinctes. Respecter refus, temporisation et mises en cache autorisées.", "Exemples : BNS non commercial ; droits de tiers FRED"],
        ["Prochaine étape", "Construire un petit pilote traçable", "Choisir quelques pays et sociétés, brancher les sources prioritaires, comparer les chiffres aux documents officiels et mesurer l’historique/complétude avant d’élargir la rotation autonome.", "Ordre proposé dans Integration"],
    ]
    append_rows(ws, rows)
    for row in range(6, ws.max_row + 1):
        ws.cell(row, 1).font = Font(name="Calibri", size=11, bold=True, color=TEAL)
        ws.cell(row, 2).font = Font(name="Calibri", size=11, bold=True, color=INK)
    return ws


def build_sources(wb, sources):
    columns = ["ID", "Source / outil", "Catégorie", "Zone couverte", "Priorité", "Preuve", "Accès gratuit et droits", "Authentification", "Données disponibles", "Historique", "Quota / cadence", "Automatisation envisagée", "Limites", "Note d’assemblage"]
    ws = prepare_sheet(wb, "Sources", "Filtrer par priorité, zone ou niveau de preuve ; toutes les conditions décrites doivent rester attachées à la source.", columns,
                       [10, 43, 38, 46, 10, 34, 72, 65, 82, 75, 78, 82, 100, 72], freeze="C6")
    ordered = sorted(sources, key=lambda item: (item["priority"], item["source"].casefold()))
    append_rows(ws, [[x["source_id"], x["source"], x.get("category"), x.get("geography"), x["priority"], EVIDENCE_LABELS[x["evidence_status"]], x.get("free_access"), x.get("auth"), x.get("data"), x.get("history"), x.get("quota"), x.get("automation"), x.get("limitations"), x["assembly_note"]] for x in ordered], "CatalogueSources")
    for row, source in enumerate(ordered, 6):
        format_priority(ws.cell(row, 5))
        ws.cell(row, 2).font = Font(name="Calibri", size=10, bold=True, color=INK)
        if source["evidence_status"] == "needs_account_validation":
            ws.cell(row, 6).fill = PatternFill("solid", fgColor=AMBER)
        elif source["evidence_status"] == "live_tested":
            ws.cell(row, 6).fill = PatternFill("solid", fgColor=PALE)
    return ws


def url_kind(url):
    lower = url.casefold()
    if any(part in lower for part in ("pricing", "prices", "terms", "tos.", "robots.txt", "/legal", "disclaimer")):
        return "Offre / conditions"
    if "github.com" in lower:
        return "Projet / code officiel"
    if any(part in lower for part in ("/api/", "openapi.", "/api?", ".json", ".csv", ".txt", ".xlsx", ".zip")):
        return "API / fichier public"
    return "Documentation / publication"


def build_links(wb, sources):
    ws = prepare_sheet(wb, "Liens", "Une URL par ligne ; les liens renvoient aux pages des fournisseurs, autorités ou projets cités. Une URL peut soutenir plusieurs fiches.",
                       ["ID source", "Source / outil", "Type de lien", "URL officielle / projet", "Portée de la preuve"], [12, 50, 30, 125, 82], freeze="C6")
    rows = []
    for item in sources:
        for url in item["urls"]:
            rows.append([item["source_id"], item["source"], url_kind(url), url,
                         "Voir Sources pour les conditions ; seuls les appels explicitement recensés dans Tests ont été testés."])
    append_rows(ws, rows, "LiensSources")
    first_rows = {}
    for row in range(6, ws.max_row + 1):
        link_cell(ws.cell(row, 4), ws.cell(row, 4).value)
        first_rows.setdefault(ws.cell(row, 1).value, row)
    source_sheet = wb["Sources"]
    for row in range(6, source_sheet.max_row + 1):
        first = first_rows.get(source_sheet.cell(row, 1).value)
        if first:
            cell = source_sheet.cell(row, 1)
            cell.hyperlink = f"#'Liens'!A{first}"
            cell.font = Font(name="Calibri", size=10, underline="single", color="176C9A")
    return ws


def build_tests(wb, tests):
    ws = prepare_sheet(wb, "Tests", "HTTP 200 indique un transport réussi ; lire la nature de la réponse et la validation métier avant de conclure à des données exploitables.",
                       ["ID", "Source / endpoint", "Portée du test", "Date / heure disponible", "HTTP mesuré", "Nature du retour", "Validation métier", "Observation résumée", "Limite du test", "Taille reçue (octets)", "URL appelée", "Fichier de preuve"],
                       [10, 48, 62, 32, 15, 44, 65, 104, 104, 20, 125, 32], freeze="C6")
    append_rows(ws, [[x["test_id"], x["source"], x["test_scope"], x["timestamp"], x["http_status"] if x["http_status"] is not None else "Non exposé", x["response_level"], x["business_validation"], x["summary"], x["limitations"], x["bytes"] if x["bytes"] is not None else "Non mesurée", x["url"], x["input_file"]] for x in tests], "TestsAcces")
    for row, item in enumerate(tests, 6):
        link_cell(ws.cell(row, 11), item["url"])
        ws.cell(row, 5).alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        ws.cell(row, 10).number_format = '#,##0'
        if item["document_only"] or "Aucune donnée" in item["response_level"] or "erreurs" in item["response_level"]:
            for col in (6, 7):
                ws.cell(row, col).fill = PatternFill("solid", fgColor=AMBER)
        else:
            ws.cell(row, 6).fill = PatternFill("solid", fgColor=PALE)
    return ws


def selected_source_names(sources, terms):
    selected = []
    for term in terms:
        for item in sources:
            if term.casefold() in item["source"].casefold() and item["source_id"] not in [x[0] for x in selected]:
                selected.append((item["source_id"], item["source"]))
    return "\n".join(f"{code} · {name}" for code, name in selected)


def build_integration(wb, sources):
    rows = [
        [1, "Identité et univers", ["Nasdaq Trader", "OpenFIGI", "GLEIF", "SEC EDGAR — comptes"], "Construire émetteur, instrument, classe, place, devise, CIK/FIGI/LEI et dates d’appartenance à l’univers.", "Distinguer entité et titre ; filtrer ETF, tests, préférentielles et warrants selon le périmètre choisi.", "Un ticker seul ou le nom de la société comme clé mondiale."],
        [2, "Comptes annuels États-Unis", ["SEC EDGAR — comptes", "SimFin"], "Fournir au moins trois exercices cohérents pour les règles existantes ; étendre l’historique si disponible.", "Rapprocher les concepts et périodes, conserver données brutes/provenance, vérifier états annuels, monnaies et versions des dépôts. Tester les droits réels SimFin.", "Utiliser fy seul ; remplacer un tag manquant par zéro ; confondre une statistique du document et une période annuelle."],
        [3, "Comptes Europe / Royaume-Uni", ["XBRL International", "DILA / AMF", "FCA National", "Companies House"], "Télécharger rapports structurés ou documents à partir des index publics, puis extraire les champs utiles.", "Valider dimensions, taxonomie, consolidé/séparé et erreurs XBRL ; vérifier par pays les droits des API et téléchargements.", "Assimiler liste de rapports, schéma API ou extraction avec erreurs à des comptes fiables et normalisés."],
        [4, "Sociétés suisses", ["Relations investisseurs", "Suisse — rapports"], "Constituer des adaptateurs d’émetteurs à partir de XLSX/XBRL/PDF officiels.", "Nestlé prouve un fichier lisible, pas tous les états de toutes les entreprises. Conserver units, exercices et liens.", "Compter le cadre SIX et un exemple Nestlé comme deux flux complets indépendants."],
        [5, "Comptes Asie et Brésil", ["EDINET", "OpenDART", "TWSE", "Brésil"], "Élargir pays par pays à des dépôts officiels et à des jeux structurés.", "Activer comptes/clés nécessaires, vérifier couverture et licences, traduire champs, calendriers locaux, unités et agrégations.", "Présenter l’instantané TWSE 2026 T2 comme plusieurs années ; présumer que les sociétés mondiales ont toutes un dépôt SEC."],
        [6, "Compléments mondiaux récents", ["ROIC.ai", "Yahoo Finance", "Fiscal.ai", "Alpha Vantage", "Financial Modeling Prep"], "Combler des champs ou vérifier une petite liste dans les limites gratuites autorisées.", "ROIC.ai : deux années ne suffisent pas seules au minimum de trois. yfinance : quatre années peuvent suffire si complètes. Mesurer couverture réelle des offres à échantillons imposés.", "Multiplier les clés ou comptes pour éviter les quotas ; confondre nombre d’endpoints et nombre de sociétés couvertes."],
        [7, "Cours et opérations sur titres", ["Yahoo Finance", "Finnhub", "Twelve Data", "Stooq"], "Obtenir prix daté, devise, splits et dividendes afin de calculer une valorisation cohérente.", "Aucune série mondiale de prix ajustés n’a été validée par les tests de cette recherche. Vérifier accès gratuit réel, ajustements, profondeur et droits d’usage.", "Utiliser des cours de démonstration comme marché complet ; estimer une capitalisation avec une moyenne diluée ou un ratio ADR erroné."],
        [8, "Dividendes et nombre d’actions", ["SEC EDGAR — comptes", "Relations investisseurs", "TWSE"], "Distinguer actions à la fin de période, moyenne diluée et nombre de couverture ; qualifier les dividendes ordinaires.", "Conserver les contextes/classes et preuves ; distinguer distributions déclarées et payées. Source générique de dividendes à vérifier contre les rapports.", "Tag absent = zéro ; tous dividendes = dividendes ordinaires ; montant par action = montant total."],
        [9, "Devises et contexte macro", ["Banque centrale européenne", "Banque nationale suisse", "FRED", "Banque mondiale", "EIA"], "Collecter les séries partagées une fois pour le bot ; conversions de référence et contexte économique.", "Documenter sens de conversion, date, moyenne ou fin de période. FRED : clé personnelle, droits de tiers et vintages ALFRED ; BNS : conditions non commerciales.", "Considérer les taux de référence comme prix d’exécution ; déduire l’exposition géographique réelle d’un groupe du seul pays de cotation."],
        [10, "Actionnariat et veille facultative", ["SEC EDGAR — Form 13F", "GDELT", "Stock Analysis"], "Ajouter des liens et alertes de contexte aux sociétés déjà filtrées.", "Respecter dates de publication, contexte des transactions et conditions de collecte ; accès StockAnalysis conditionnel aux CGU/robots.txt, sans API confirmée.", "Copier aveuglément les 13F différés ; confondre achat de marché et attribution ; utiliser le bruit médiatique comme preuve de qualité financière."],
        [11, "Mesure de couverture et qualité", ["SEC EDGAR — comptes", "XBRL International", "OpenFIGI"], "Mesurer société par société les champs, périodes et sources disponibles avant tout classement.", "Conserver provenance par mesure, qualifier conflits, lever l’abstention seulement si les données suffisent. Séparer familles de phases et informations sectorielles.", "Promettre une recherche mondiale exhaustive ou transformer un score heuristique en probabilité de gain."],
        [12, "Automatisation du pilote", ["SEC EDGAR — comptes", "Nasdaq Trader", "Banque centrale européenne"], "Planifier des mises à jour incrémentales après validation d’un univers limité et de ses sources.", "Cache, historique, quotas par fournisseur, arrêt sur refus, rapport de couverture réelle et journal des échecs. Les connecteurs restent à construire et à valider.", "Qualifier cette recherche documentaire de déploiement autonome déjà actif."],
    ]
    values = [[order, need, selected_source_names(sources, terms), target, prerequisites, avoid, "Proposé — collecteurs à construire / valider"] for order, need, terms, target, prerequisites, avoid in rows]
    ws = prepare_sheet(wb, "Integration", "Ordre de travail proposé : d’abord des identités et des comptes fiables, puis les prix, la couverture et l’automatisation. Aucun achat de marché n’est déclenché.",
                       ["Ordre", "Besoin du bot", "Sources / outils associés", "Objectif", "Conditions avant automatisation", "Erreur à éviter", "État"],
                       [10, 36, 74, 86, 105, 95, 44], freeze="C6")
    append_rows(ws, values, "IntegrationSources")
    for row in range(6, ws.max_row + 1):
        ws.cell(row, 1).font = Font(name="Calibri", size=14, bold=True, color=TEAL)
        ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="top")
        ws.cell(row, 2).font = Font(name="Calibri", size=11, bold=True, color=INK)
    return ws


def validate_workbook(path, source_count, link_count, test_count):
    wb = load_workbook(path, read_only=False, data_only=False)
    expected = ["Synthese", "Sources", "Liens", "Tests", "Integration"]
    if wb.sheetnames != expected:
        raise AssertionError(f"Onglets inattendus : {wb.sheetnames}")
    for sheet, expected_count in [("Sources", source_count), ("Liens", link_count), ("Tests", test_count), ("Integration", 12)]:
        if wb[sheet].max_row - 5 != expected_count:
            raise AssertionError(f"Nombre de lignes inattendu dans {sheet}")
        if wb[sheet].freeze_panes is None or not wb[sheet].auto_filter.ref:
            raise AssertionError(f"Filtres/volets absents dans {sheet}")
    formula_cells = []
    for ws in wb:
        for merged in ws.merged_cells.ranges:
            if merged.max_row >= 5:
                raise AssertionError(f"Fusion de cellules de données : {ws.title}!{merged}")
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f":
                    formula_cells.append(f"{ws.title}!{cell.coordinate}")
    if formula_cells:
        raise AssertionError(f"Formules inattendues : {formula_cells}")
    links = [row[3].hyperlink for row in wb["Liens"].iter_rows(min_row=6)]
    if len(links) != link_count or not all(links):
        raise AssertionError("Hyperliens manquants")
    source_rows = list(wb["Sources"].iter_rows(min_row=6, values_only=True))
    for row in source_rows:
        if "yfinance" in str(row[1]).casefold() or "fiscal.ai" in str(row[1]).casefold():
            if row[4] != "B":
                raise AssertionError("Rétrogradation B absente")
    wb.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=BASE)
    parser.add_argument("--output-dir", type=Path, default=BASE.parent)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    sources = load_sources(args.source_dir)
    originals, tests, manifests = load_tests(args.source_dir)
    all_urls = [url for item in sources for url in item["urls"]]
    stats = {
        "source_and_tool_cards": len(sources),
        "priority_counts": dict(Counter(item["priority"] for item in sources)),
        "evidence_counts": dict(Counter(item["evidence_status"] for item in sources)),
        "source_url_rows": len(all_urls), "unique_source_urls": len(set(all_urls)),
        "test_observations": len(tests),
        "http_200_count": sum(item["http_status"] == 200 for item in tests),
        "http_429_count": sum(item["http_status"] == 429 for item in tests),
        "http_status_unexposed_count": sum(item["http_status"] is None for item in tests),
        "returned_data_or_metadata_count": sum(item["returned_data_or_metadata"] for item in tests),
        "document_only_count": sum(item["document_only"] for item in tests),
        "no_data_observations_count": sum("Aucune donnée" in item["response_level"] for item in tests),
        "new_collectors_integrated_into_existing_bot": 0,
    }
    now = datetime.now(timezone.utc).isoformat()
    evidence = {
        "generated_at_utc": now,
        "scope": "Catalogue de recherche et observations d’accès. Aucun scan mondial exhaustif validé ; aucun nouveau collecteur gratuit intégré au bot par cette recherche.",
        "interpretation": "HTTP 200 ≠ données financières utilisables. Distinguer documents techniques, métadonnées, faits bruts et validation métier. Les codes et heures non observés ne sont pas inventés.",
        "summary": stats,
        "input_manifests": manifests,
        "original_evidence_batches": originals,
        "normalized_test_observations": tests,
        "source_documents": [{"source_id": x["source_id"], "source": x["source"], "catalogue_file": x["catalogue_file"], "priority": x["priority"], "priority_original": x["priority_original"], "assembly_note": x["assembly_note"], "urls": x["urls"]} for x in sources],
    }
    evidence_path = args.output_dir / "Tests_acces_donnees_gratuites.json"
    evidence_path.write_text(json.dumps(evidence, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "Recherche de données boursières"
    wb.properties.title = "Catalogue des données boursières gratuites — recherche au 6 octobre 2026"
    wb.properties.subject = "Sources et outils, conditions gratuites, preuves d’accès et intégration proposée"
    wb.properties.description = "Données publiques et offres gratuites examinées ; couverture réelle et collecteurs à valider."
    build_summary(wb, sources, tests, stats)
    build_sources(wb, sources)
    build_links(wb, sources)
    build_tests(wb, tests)
    build_integration(wb, sources)
    wb.active = 0
    output = args.output_dir / "Catalogue_donnees_boursieres_gratuites.xlsx"
    wb.save(output)
    wb.close()
    validate_workbook(output, len(sources), len(all_urls), len(tests))
    parsed_evidence = read_json(evidence_path)
    if parsed_evidence["original_evidence_batches"] != originals:
        raise AssertionError("Altération des preuves d’origine")
    print(json.dumps({"xlsx": str(output.resolve()), "evidence_json": str(evidence_path.resolve()), "validation": "Lecture XLSX/JSON, onglets, comptes, filtres, volets, hyperliens, absence de formules et de fusions de données : OK", "summary": stats}, ensure_ascii=False))


if __name__ == "__main__":
    main()
